#!/usr/bin/env python3
"""
Generate a Word document consolidating iperf Network Throughput Test results.
Reads iperf_*_result.log files from testresult/ and produces a structured
.docx report with: Title, Summary table, TCP details, Server logs, Conclusion.
"""

import os
import sys
import glob
import re
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

COLOR_GREEN = RGBColor(0x1A, 0x7F, 0x37)
COLOR_RED = RGBColor(0xD1, 0x24, 0x2F)
COLOR_GREY = RGBColor(100, 100, 100)

BW_RE = re.compile(
    r"\[\s*\d+\]\s+[\d.]+-[\d.]+\s+sec\s+"
    r"([\d.]+\s+\S+)\s+([\d.]+\s+\S+)"
)


def find_iperf_logs(base_dir):
    pattern = os.path.join(base_dir, "iperf_*_result.log")
    return [f for f in sorted(glob.glob(pattern))
            if not os.path.basename(f).startswith("iperf_srv_")]


def find_iperf_server_logs(base_dir):
    return sorted(glob.glob(os.path.join(base_dir, "iperf_srv_*")))



def parse_iperf_log(filepath):
    """Parse an iperf client result log and extract structured data."""
    with open(filepath, "r") as f:
        content = f.read()

    d = dict(filename=os.path.basename(filepath), hostname="", test_date="",
             aix_version="", iperf_cmd="", server_ip="", port="", duration="",
             tcp_bandwidth="", tcp_transfer="", tcp_window="", tcp_raw="",
             udp_bandwidth="", udp_transfer="", udp_datagrams="",
             udp_target_bw="", udp_raw="", raw=content)

    for line in content.splitlines():
        s = line.strip()
        if s.startswith("Hostname:"):
            d["hostname"] = s.split(":", 1)[1].strip()
        elif s.startswith("Test Date:"):
            d["test_date"] = s.split(":", 1)[1].strip()
        elif s.startswith("AIX Version:"):
            d["aix_version"] = s.split(":", 1)[1].strip()
        elif s.startswith("iperf Command:"):
            d["iperf_cmd"] = s.split(":", 1)[1].strip()
        elif "[INFO] Target Server:" in s:
            d["server_ip"] = s.split(":", 1)[1].strip()
        elif "[INFO] Port:" in s:
            d["port"] = s.split(":", 1)[1].strip()
        elif "[INFO] Test Duration:" in s:
            d["duration"] = s.split(":", 1)[1].strip()

    # TCP results
    tcp_sec = re.search(
        r"TCP Throughput Test(.*?)(?=UDP Throughput Test|Network Throughput Tests Complete)",
        content, re.DOTALL)
    if tcp_sec:
        d["tcp_raw"] = tcp_sec.group(0).strip()
        m = BW_RE.search(tcp_sec.group(1))
        if m:
            d["tcp_transfer"] = m.group(1)
            d["tcp_bandwidth"] = m.group(2)

    tw = re.search(r"TCP window size:\s+(.+)", content)
    if tw:
        d["tcp_window"] = tw.group(1).strip()

    # UDP results
    udp_sec = re.search(
        r"UDP Throughput Test(.*?)(?=Network Throughput Tests Complete)",
        content, re.DOTALL)
    if udp_sec:
        d["udp_raw"] = udp_sec.group(0).strip()
        m = BW_RE.search(udp_sec.group(1))
        if m:
            d["udp_transfer"] = m.group(1)
            d["udp_bandwidth"] = m.group(2)

    dg = re.search(r"Sent\s+(\d+)\s+datagrams", content)
    if dg:
        d["udp_datagrams"] = dg.group(1)

    ut = re.search(r"\[INFO\]\s+Testing UDP throughput with bandwidth:\s+(\S+)", content)
    if ut:
        d["udp_target_bw"] = ut.group(1)

    return d


def parse_server_log(filepath):
    """Parse a server-side iperf log."""
    with open(filepath, "r") as f:
        content = f.read()
    d = dict(filename=os.path.basename(filepath), raw=content,
             bandwidth="", transfer="")
    m = BW_RE.search(content)
    if m:
        d["transfer"] = m.group(1)
        d["bandwidth"] = m.group(2)
    return d



def _bold_cell(cell, text):
    """Set cell text with bold formatting."""
    cell.text = text
    for p in cell.paragraphs:
        for r in p.runs:
            r.bold = True


def _kv_table(doc, rows):
    """Add a 2-column key-value table."""
    tbl = doc.add_table(rows=len(rows), cols=2)
    tbl.style = "Light Grid Accent 1"
    for i, (label, val) in enumerate(rows):
        _bold_cell(tbl.rows[i].cells[0], label)
        tbl.rows[i].cells[1].text = val or "\u2014"
    return tbl


def _raw_block(doc, text):
    """Add a monospace raw-output block."""
    p = doc.add_paragraph()
    run = p.add_run("Raw Output:")
    run.bold = True
    rp = doc.add_paragraph()
    rr = rp.add_run(text)
    rr.font.size = Pt(8)
    rr.font.name = "Consolas"


def generate_iperf_report(base_dir, output_path,
                          client_files=None, server_files=None):
    """Generate a consolidated iperf report .docx.
    When client_files/server_files are provided, only those files are used.
    Otherwise, auto-discover all iperf logs in base_dir."""
    client_logs = client_files if client_files else find_iperf_logs(base_dir)
    server_logs = server_files if server_files is not None else find_iperf_server_logs(base_dir)

    if not client_logs:
        print("No iperf client result logs found.")
        return False

    all_client = [parse_iperf_log(f) for f in client_logs]
    all_server = [parse_server_log(f) for f in server_logs]

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    # Tighten heading styles: reduce vertical spacing and add left indent
    for lvl, (sp_before, sp_after, indent) in {
        1: (Pt(12), Pt(4), Cm(0)),
        2: (Pt(8),  Pt(3), Cm(0.5)),
        3: (Pt(6),  Pt(2), Cm(1.0)),
    }.items():
        hstyle = doc.styles[f"Heading {lvl}"]
        hstyle.paragraph_format.space_before = sp_before
        hstyle.paragraph_format.space_after = sp_after
        hstyle.paragraph_format.left_indent = indent

    # ── Title Page ─────────────────────────────────────
    doc.add_paragraph()
    title = doc.add_heading("Network Throughput Test Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    hostnames = [d["hostname"] or "Unknown" for d in all_client]
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(
        f"\nClient Host(s): {', '.join(hostnames)}\n"
        f"Test Count: {len(all_client)}\n")
    run.font.size = Pt(14)

    gen = doc.add_paragraph()
    gen.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = gen.add_run(
        f"Report Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_GREY
    doc.add_page_break()

    # ── 1. Executive Summary ───────────────────────────
    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "This report consolidates iperf network throughput tests "
        "between AIX LPARs. Each test measures TCP bandwidth "
        "between a client and server LPAR.")

    cols = ["#", "Client Host", "Server IP", "Test Date",
            "TCP Bandwidth"]
    tbl = doc.add_table(rows=1 + len(all_client), cols=len(cols))
    tbl.style = "Light Grid Accent 1"
    for i, h in enumerate(cols):
        _bold_cell(tbl.rows[0].cells[i], h)

    for idx, d in enumerate(all_client):
        row = tbl.rows[idx + 1]
        row.cells[0].text = str(idx + 1)
        row.cells[1].text = d["hostname"] or "\u2014"
        row.cells[2].text = d["server_ip"] or "\u2014"
        row.cells[3].text = d["test_date"] or "\u2014"
        row.cells[4].text = d["tcp_bandwidth"] or "\u2014"

    doc.add_paragraph()
    doc.add_page_break()


    # ── 2. Per-Test Details ────────────────────────────
    doc.add_heading("2. Test Details", level=1)

    for idx, d in enumerate(all_client):
        doc.add_heading(
            f"2.{idx+1} {d['hostname'] or 'Test ' + str(idx+1)}", level=2)

        _kv_table(doc, [
            ("Client Hostname", d["hostname"]),
            ("AIX Version", d["aix_version"]),
            ("Target Server IP", d["server_ip"]),
            ("iperf Port", d["port"]),
            ("Test Duration", d["duration"]),
            ("iperf Command", d["iperf_cmd"]),
            ("Test Date", d["test_date"]),
        ])
        doc.add_paragraph()

        # TCP
        doc.add_heading("TCP Throughput", level=3)
        _kv_table(doc, [
            ("TCP Window Size", d["tcp_window"]),
            ("Transfer", d["tcp_transfer"]),
            ("Bandwidth", d["tcp_bandwidth"]),
        ])
        if d["tcp_raw"]:
            doc.add_paragraph()
            _raw_block(doc, d["tcp_raw"])
        doc.add_paragraph()

    # ── 3. Server-Side Logs ────────────────────────────
    if all_server:
        doc.add_page_break()
        doc.add_heading("3. Server-Side Logs", level=1)
        doc.add_paragraph(
            "The following data was captured on the iperf server side.")
        for sd in all_server:
            doc.add_heading(sd["filename"], level=3)
            if sd["bandwidth"]:
                p = doc.add_paragraph()
                p.add_run("Server-observed bandwidth: ").bold = True
                p.add_run(
                    f"{sd['bandwidth']}  (Transfer: {sd['transfer']})")
            rp = doc.add_paragraph()
            rr = rp.add_run(sd["raw"])
            rr.font.size = Pt(8)
            rr.font.name = "Consolas"
            doc.add_paragraph()

    # ── 4. Conclusion ──────────────────────────────────
    doc.add_page_break()
    sec = 4 if all_server else 3
    doc.add_heading(f"{sec}. Conclusion", level=1)

    all_tcp = all(d["tcp_bandwidth"] for d in all_client)

    if all_tcp:
        para = doc.add_paragraph()
        run = para.add_run("PASS")
        run.bold = True
        run.font.color.rgb = COLOR_GREEN
        para.add_run(
            " \u2014 All iperf tests completed successfully. "
            "TCP throughput measurements were captured "
            "for all tested LPAR pairs.")
    else:
        para = doc.add_paragraph()
        run = para.add_run("REVIEW")
        run.bold = True
        run.font.color.rgb = COLOR_RED
        para.add_run(
            " \u2014 Some tests may have incomplete results. "
            "Review the per-test details above.")

    for idx, d in enumerate(all_client):
        ok = bool(d["tcp_bandwidth"])
        para = doc.add_paragraph(style="List Bullet")
        sr = para.add_run("\u2713 PASS" if ok else "\u26A0 REVIEW")
        sr.bold = True
        sr.font.color.rgb = COLOR_GREEN if ok else COLOR_RED
        para.add_run(
            f" \u2014 {d['hostname'] or 'Test ' + str(idx+1)} "
            f"\u2192 {d['server_ip'] or '?'}: "
            f"TCP {d['tcp_bandwidth'] or 'N/A'}")

    # ── Appendix: Full Test Logs ───────────────────────
    doc.add_page_break()
    app_sec = sec + 1
    doc.add_heading(f"{app_sec}. Appendix: Full Test Logs", level=1)
    doc.add_paragraph(
        "This appendix contains the complete, unmodified log files "
        "captured during each iperf test for reference and audit purposes.")

    for idx, d in enumerate(all_client):
        doc.add_heading(
            f"{app_sec}.{idx+1} Client Log \u2014 "
            f"{d['hostname'] or 'Unknown'} ({d['filename']})",
            level=2)
        lp = doc.add_paragraph()
        lr = lp.add_run(d["raw"])
        lr.font.size = Pt(7)
        lr.font.name = "Consolas"
        lr.font.color.rgb = COLOR_GREY

    if all_server:
        for idx, sd in enumerate(all_server):
            doc.add_heading(
                f"{app_sec}.{len(all_client)+idx+1} Server Log \u2014 "
                f"{sd['filename']}",
                level=2)
            lp = doc.add_paragraph()
            lr = lp.add_run(sd["raw"])
            lr.font.size = Pt(7)
            lr.font.name = "Consolas"
            lr.font.color.rgb = COLOR_GREY

    # Save
    try:
        doc.save(output_path)
        print(f"Report generated: {output_path}")
    except PermissionError:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        root, ext = os.path.splitext(output_path)
        fallback = f"{root}_{ts}{ext}"
        print(f"WARNING: Cannot write '{output_path}' (file open?).")
        doc.save(fallback)
        print(f"Report generated: {fallback}")

    return True


def main():
    base_dir = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "testresult")
    if not os.path.exists(base_dir):
        print(f"Error: testresult directory not found at {base_dir}")
        sys.exit(1)

    # Parse --files <name1> <name2> ... to filter which logs to include
    selected_files = None
    if "--files" in sys.argv:
        idx = sys.argv.index("--files")
        selected_files = [os.path.basename(f) for f in sys.argv[idx + 1:]]

    if selected_files:
        # Only use the explicitly listed files (resolve basenames in base_dir)
        logs = []
        srv_logs = []
        for name in selected_files:
            full = os.path.join(base_dir, name)
            if not os.path.isfile(full):
                continue
            if name.startswith("iperf_srv_"):
                srv_logs.append(full)
            elif "iperf" in name and name.endswith("_result.log"):
                logs.append(full)
            elif "iperf" in name and name.endswith(".log"):
                srv_logs.append(full)
    else:
        logs = find_iperf_logs(base_dir)
        srv_logs = None  # let generate_iperf_report auto-discover

    if not logs:
        print("Error: No iperf result logs found in testresult/")
        sys.exit(1)

    print(f"Found {len(logs)} iperf result log(s):")
    for f in logs:
        print(f"  - {os.path.basename(f)}")

    output_path = os.path.join(base_dir, "Network_Throughput_Report.docx")
    print("\nGenerating iperf report...")
    generate_iperf_report(base_dir, output_path,
                          client_files=logs, server_files=srv_logs)


if __name__ == "__main__":
    main()

